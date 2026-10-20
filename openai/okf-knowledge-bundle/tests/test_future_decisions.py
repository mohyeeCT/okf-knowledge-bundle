"""Future-dated decisions fail at intake without changing saved review state."""
import copy
from datetime import datetime, timezone
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from okf_bundle.cli import execute, parser
from okf_bundle.review import accept_availability, apply_review, approve_alias, import_drafts, reviewer_event
from okf_bundle.state import StateStore
from okf_bundle.workbook import export_review, import_review
from test_export import fixture_state
from test_review_state import draft_state

CREATED = "2026-10-15T10:00:00Z"
CURRENT = "2026-10-15T12:00:00.500000Z"
FUTURE = "2026-10-15T12:00:00.500001Z"


class FutureDecisionTests(unittest.TestCase):
    def setUp(self):
        clock_patch = patch("okf_bundle.review.datetime", wraps=datetime)
        clock = clock_patch.start()
        self.addCleanup(clock_patch.stop)
        clock.now.return_value = datetime(2026, 10, 15, 12, 0, 0, 500000, tzinfo=timezone.utc)

    def drafted(self, risk=False):
        state, source, payload = draft_state()
        if risk:
            payload["claims"][0]["risk_categories"] = ["price"]
        with patch("okf_bundle.review.utc_now", return_value=CREATED):
            self.assertTrue(import_drafts(state, payload)["accepted"])
        return state, state["sources"][source["id"]]

    def test_event_rejects_future_and_preserves_current_fractional_or_offset_time(self):
        for at in (CURRENT, "2026-10-15T08:00:00.500000-04:00", CREATED):
            with self.subTest(at=at):
                self.assertEqual(at, reviewer_event("human:analyst", at, "synthetic", "fixture", "revision")["at"])
        with self.assertRaisesRegex(ValueError, "future"):
            reviewer_event("human:analyst", FUTURE, "synthetic", "fixture", "revision")

    def test_claim_and_concept_future_decisions_are_atomic(self):
        cases = [("claim", {"evidence_review": "accepted"}, "human:analyst"),
                 ("claim", {"publication": "approved"}, "human:analyst"),
                 ("claim", {"business_confirmation": "confirmed"}, "human:client-owner"),
                 ("concept", {"decision": "accepted"}, "human:analyst")]
        for entity, decisions, reviewer in cases:
            with self.subTest(entity=entity, decisions=decisions):
                state, _ = self.drafted()
                record = next(iter(state[entity + "s"].values()))
                before = copy.deepcopy(state)
                with self.assertRaisesRegex(ValueError, "future"):
                    apply_review(state, entity, record["id"], record["revision"], decisions, reviewer,
                                 FUTURE, "synthetic", "not-real-approval")
                self.assertEqual(before, state)

    def test_alias_and_availability_reject_future_without_mutation(self):
        state, source = self.drafted()
        target = copy.deepcopy(source)
        target.update(id="alias-target", source_key="alias-target", requested_url="https://example.com/alias/",
                      normalized_url="https://example.com/alias/", citation_url="https://example.com/alias/",
                      confirmed_identity_url="https://example.com/alias/")
        state["sources"][target["id"]] = target
        before = copy.deepcopy(state)
        with self.assertRaisesRegex(ValueError, "future"):
            approve_alias(state, source["id"], target["id"], "human:analyst", FUTURE,
                          "synthetic", "not-real-approval", "Synthetic equal evidence review")
        self.assertEqual(before, state)
        with self.assertRaisesRegex(ValueError, "future"):
            accept_availability(state, [source["id"]], "human:client-owner", FUTURE,
                                "Synthetic unknown availability acceptance", "synthetic", "not-real-approval")
        self.assertEqual(before, state)
        accept_availability(state, [source["id"]], "human:client-owner", CURRENT,
                            "Synthetic unknown availability acceptance", "synthetic", "not-real-approval")
        self.assertEqual(CURRENT, source["availability_acceptance"]["at"])
        approve_alias(state, source["id"], target["id"], "human:analyst", CURRENT,
                      "synthetic", "not-real-approval", "Synthetic equal evidence review")
        self.assertEqual(target["id"], source["alias_of"])

    def test_future_workbook_cannot_partially_apply_claim_concept_or_scope_changes(self):
        state, _ = self.drafted()
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "review.xlsx"
            export_review(state, path)
            book = load_workbook(path)
            book["Claims"]["G2"] = "accepted"
            book["Concepts"]["H2"] = "accepted"
            book["Classification"]["H2"] = "excluded"
            book["Classification"]["I2"] = "Synthetic scope review"
            book.save(path)
            book.close()
            before = copy.deepcopy(state)
            with self.assertRaisesRegex(ValueError, "future"):
                import_review(state, path, "human:analyst", FUTURE, "synthetic", "not-real-approval")
            self.assertEqual(before, state)
            self.assertEqual(3, import_review(state, path, "human:analyst", CURRENT,
                                            "synthetic", "not-real-approval")["changed_rows"])

    def test_future_client_workbook_does_not_store_confirmation_or_private_correction(self):
        state, _ = self.drafted(risk=True)
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "client.xlsx"
            export_review(state, Path(folder) / "review.xlsx", path)
            book = load_workbook(path)
            book["Client confirmation"]["E2"] = "no"
            book["Client confirmation"]["F2"] = "Synthetic correction requiring public support"
            book["Client confirmation"]["G2"] = "Fictional client scope"
            book.save(path)
            book.close()
            before = copy.deepcopy(state)
            with self.assertRaisesRegex(ValueError, "future"):
                import_review(state, path, "human:client-owner", FUTURE, "synthetic", "not-real-approval", client=True)
            self.assertEqual(before, state)

    def test_future_source_review_and_release_approval_leave_disk_state_unchanged(self):
        with tempfile.TemporaryDirectory() as folder:
            store = StateStore(Path(folder) / "run")
            state = fixture_state()
            with store.transaction():
                store.save(state)
            shared = ["--run", str(store.root)]
            reviewer = ["--by", "human:client-owner", "--at", FUTURE, "--channel", "synthetic", "--reference", "not-real-approval"]
            before = store.load()
            with self.assertRaisesRegex(ValueError, "future"):
                execute(parser().parse_args(["source-review", *shared, "--id", "source-1",
                    "--url", "https://example.com/training/", "--scope", "included", "--reason", "Synthetic source review", *reviewer]))
            self.assertEqual(before, store.load())
            candidate = execute(parser().parse_args(["candidate", *shared, "--release-id", "time-bound", "--as-of", CREATED]))
            before = store.load()
            with self.assertRaisesRegex(ValueError, "future"):
                execute(parser().parse_args(["release-approve", *shared, "--candidate", candidate["candidate"],
                    "--manifest-hash", candidate["manifest_hash"], *reviewer]))
            self.assertEqual(before, store.load())


if __name__ == "__main__":
    unittest.main()
