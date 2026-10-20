"""Behavioral tests of human-review boundaries and transactional state."""
import copy
from datetime import datetime, timezone
import json
import os
from pathlib import Path
import tempfile
import unittest
import zipfile
from unittest.mock import patch
from openpyxl import load_workbook
from okf_bundle.common import digest, normalize_url, utc_now
from okf_bundle.state import StateStore, new_state
from okf_bundle.ingest import new_source, add_version
from okf_bundle.extraction import extract_html
from okf_bundle.review import import_drafts, apply_review, accept_availability, approve_alias
from okf_bundle.workbook import export_review, import_review
from okf_bundle.workflow import merge_import, packets
from okf_bundle.availability import check_availability


def draft_state():
    state = new_state({"domains": ["example.com"], "publication_owner": "human:client-owner", "confirmation_policy": "all-risk"})
    source = new_source("https://example.com/training/", state["config"])
    raw = '<html lang="en"><head><title>Training</title></head><body><main><h1>Team training</h1><p>Training is offered on site.</p><p>The fee is $100 per team.</p></main></body></html>'
    add_version(source, raw, extract_html(raw, source["requested_url"]), "sf-original", state["config"])
    state["sources"][source["id"]] = source
    state["inventory"] = [{"input_ref": "fixture.html", "url": source["requested_url"], "source_id": source["id"], "state": "included", "reason": "fixture"}]
    payload = {"batch_id": "fixture-1", "claims": [{"key": "onsite", "subject": "training", "predicate": "delivery", "value": "on site", "qualifiers": {}, "public_statement": "Training is offered on site.", "evidence": [{"source_id": source["id"], "version_id": source["active_version"], "locator": "body", "quote": "Training is offered on site."}]}], "concepts": [{"key": "team-training", "type": "Service", "title": "Team training", "description": "Delivery of team training.", "locale": "en", "claim_keys": ["onsite"]}]}
    return state, source, payload


class ReviewStateTests(unittest.TestCase):
    def setUp(self):
        # Fixed chronology fixtures may review on Oct 16; their clock is explicit.
        clock_patch = patch("okf_bundle.review.datetime", wraps=datetime)
        clock = clock_patch.start()
        self.addCleanup(clock_patch.stop)
        fixture_end = datetime(2026, 10, 17, tzinfo=timezone.utc)
        clock.now.side_effect = lambda tz: max(datetime.now(tz), fixture_end)

    def test_material_creation_times_survive_unchanged_reimport(self):
        state, _, payload = draft_state()
        with patch("okf_bundle.review.utc_now", return_value="2026-10-15T00:00:00Z"):
            self.assertTrue(import_drafts(state, payload)["accepted"])
        claim = next(iter(state["claims"].values()))
        concept = next(iter(state["concepts"].values()))
        with patch("okf_bundle.review.utc_now", return_value="2026-10-16T00:00:00Z"):
            self.assertFalse(import_drafts(state, payload)["material_changed"])
        self.assertEqual(claim["material_revision_created_at"], "2026-10-15T00:00:00Z")
        self.assertEqual(concept["material_revision_created_at"], "2026-10-15T00:00:00Z")

    def test_changed_claim_updates_own_and_dependent_concept_creation(self):
        state, _, payload = draft_state()
        with patch("okf_bundle.review.utc_now", return_value="2026-10-15T00:00:00Z"):
            import_drafts(state, payload)
        original_claim = next(iter(state["claims"].values()))
        old_revision = original_claim["revision"]
        concept_id = next(iter(state["concepts"]))
        apply_review(state, "concept", concept_id, state["concepts"][concept_id]["revision"], {"decision": "accepted"},
                     "human:agency-reviewer", "2026-10-15T00:00:01Z", "synthetic-test", "concept-reviewed")
        payload["claims"][0]["public_statement"] = "On-site training is offered."
        payload["concepts"] = []  # Existing concepts must invalidate even when omitted from this batch.
        with patch("okf_bundle.review.utc_now", return_value="2026-10-16T00:00:00Z"):
            import_drafts(state, payload)
        changed_claim = next(iter(state["claims"].values()))
        self.assertNotEqual(changed_claim["revision"], old_revision)
        self.assertEqual(changed_claim["material_revision_created_at"], "2026-10-16T00:00:00Z")
        self.assertEqual(state["concepts"][concept_id]["material_revision_created_at"], "2026-10-16T00:00:00Z")
        self.assertEqual(state["concepts"][concept_id]["review"]["decision"], "pending")

    def test_review_cannot_predate_claim_or_concept_revision(self):
        for entity, decisions in (("claim", {"evidence_review": "accepted"}), ("concept", {"decision": "accepted"})):
            with self.subTest(entity=entity):
                state, _, payload = draft_state()
                with patch("okf_bundle.review.utc_now", return_value="2026-10-15T00:00:00Z"):
                    import_drafts(state, payload)
                record = next(iter(state[entity + "s"].values()))
                before = copy.deepcopy(state)
                with self.assertRaisesRegex(ValueError, "predates"):
                    apply_review(state, entity, record["id"], record["revision"], decisions, "human:agency-reviewer",
                                 "2026-09-15T00:00:00Z", "synthetic-test", "backdated-review")
                self.assertEqual(state, before)

    def test_creation_boundary_accepts_equivalent_offset_time(self):
        state, _, payload = draft_state()
        with patch("okf_bundle.review.utc_now", return_value="2026-10-15T00:00:00Z"):
            import_drafts(state, payload)
        claim = next(iter(state["claims"].values()))
        apply_review(state, "claim", claim["id"], claim["revision"], {"evidence_review": "accepted"},
                     "human:agency-reviewer", "2026-10-14T20:00:00-04:00", "synthetic-test", "same-instant")
        self.assertEqual(claim["evidence_review"], "accepted")

    def test_legacy_reviews_reset_at_current_tool_boundary(self):
        state, _, payload = draft_state()
        with patch("okf_bundle.review.utc_now", return_value="2026-10-15T00:00:00Z"):
            import_drafts(state, payload)
        claim = next(iter(state["claims"].values()))
        concept = next(iter(state["concepts"].values()))
        apply_review(state, "claim", claim["id"], claim["revision"], {"evidence_review": "accepted"},
                     "human:agency-reviewer", "2026-10-15T00:00:01Z", "synthetic-test", "legacy-claim")
        apply_review(state, "concept", concept["id"], concept["revision"], {"decision": "accepted"},
                     "human:agency-reviewer", "2026-10-15T00:00:01Z", "synthetic-test", "legacy-concept")
        del claim["material_revision_created_at"]
        del concept["material_revision_created_at"]
        with self.assertRaisesRegex(ValueError, "creation time"):
            apply_review(state, "claim", claim["id"], claim["revision"], {"publication": "approved"},
                         "human:agency-reviewer", "2026-10-16T00:00:00Z", "synthetic-test", "untouched-legacy")
        state["release_approvals"] = {"synthetic-old-manifest": {"status": "test-only"}}
        with patch("okf_bundle.review.utc_now", return_value="2026-10-16T00:00:00Z"):
            self.assertTrue(import_drafts(state, payload)["material_changed"])
        current_claim = state["claims"][claim["id"]]
        current_concept = state["concepts"][concept["id"]]
        self.assertEqual(current_claim["material_revision_created_at"], "2026-10-16T00:00:00Z")
        self.assertEqual(current_claim["evidence_review"], "pending")
        self.assertEqual(current_claim["reviews"], [])
        self.assertEqual(current_claim["previous_revisions"][-1]["reviews"][0]["at"], "2026-10-15T00:00:01Z")
        self.assertIsNone(current_claim["previous_revisions"][-1]["material_revision_created_at"])
        self.assertEqual(current_concept["material_revision_created_at"], "2026-10-16T00:00:00Z")
        self.assertEqual(current_concept["review"]["decision"], "pending")
        self.assertEqual(current_concept["legacy_review_history"][-1]["decision"], "accepted")
        self.assertEqual(state["release_approvals"], {})

    def test_legacy_claim_migration_resets_modern_dependent_concept(self):
        state, _, payload = draft_state()
        with patch("okf_bundle.review.utc_now", return_value="2026-10-15T00:00:00Z"):
            import_drafts(state, payload)
        claim = next(iter(state["claims"].values()))
        concept = next(iter(state["concepts"].values()))
        apply_review(state, "concept", concept["id"], concept["revision"], {"decision": "accepted"},
                     "human:agency-reviewer", "2026-10-15T00:00:01Z", "synthetic-test", "concept-reviewed")
        del claim["material_revision_created_at"]
        payload["concepts"] = []
        with patch("okf_bundle.review.utc_now", return_value="2026-10-16T00:00:00Z"):
            import_drafts(state, payload)
        current = state["concepts"][concept["id"]]
        self.assertEqual(current["review"]["decision"], "pending")
        self.assertEqual(current["material_revision_created_at"], "2026-10-16T00:00:00Z")

    def test_model_cannot_supply_creation_time_or_resolution_reason(self):
        for field, value in (("material_revision_created_at", "2020-01-01T00:00:00Z"), ("reason", "Model declares conflict resolved")):
            with self.subTest(field=field):
                state, _, payload = draft_state()
                payload["claims"][0][field] = value
                self.assertFalse(import_drafts(state, payload)["accepted"])
                self.assertEqual(state["claims"], {})

    def test_open_conflict_needs_audited_human_resolution(self):
        state, _, payload = draft_state()
        payload["claims"][0]["risk_categories"] = ["disputed"]
        import_drafts(state, payload)
        claim = next(iter(state["claims"].values()))
        before = copy.deepcopy(state)
        for decision, reason in (("none", None), ("resolved", None), ("resolved", ""), ("resolved", "unsafe\nreason")):
            with self.subTest(decision=decision, reason=reason):
                with self.assertRaises(ValueError):
                    apply_review(state, "claim", claim["id"], claim["revision"], {"conflict": decision},
                                 "human:agency-reviewer", utc_now(), "synthetic-test", "invalid-resolution", reason=reason)
                self.assertEqual(state, before)
        apply_review(state, "claim", claim["id"], claim["revision"], {"conflict": "resolved"},
                     "human:agency-reviewer", utc_now(), "synthetic-test", "recorded-resolution",
                     reason="Both public pages now agree; updated evidence reviewed under this revision")
        event = claim["reviews"][-1]
        self.assertEqual(claim["conflict"], "resolved")
        self.assertEqual(event["dimension"], "conflict")
        self.assertEqual(event["decision"], "resolved")
        self.assertEqual(event["revision"], claim["revision"])
        self.assertEqual(event["by"], "human:agency-reviewer")
        self.assertIn("Both public pages", event["reason"])
        self.assertEqual(state["review_history"][-1]["events"][-1], event)
        with self.assertRaises(ValueError):
            apply_review(state, "claim", claim["id"], claim["revision"], {"conflict": "none"},
                         "human:agency-reviewer", utc_now(), "synthetic-test", "erase-history")

    def test_no_conflict_cannot_gain_a_fabricated_resolution(self):
        state, _, payload = draft_state()
        import_drafts(state, payload)
        claim = next(iter(state["claims"].values()))
        with self.assertRaisesRegex(ValueError, "existing open conflict"):
            apply_review(state, "claim", claim["id"], claim["revision"], {"conflict": "resolved"},
                         "human:agency-reviewer", utc_now(), "synthetic-test", "fake-resolution", reason="No conflict was opened")

    def test_singleline_draft_fields_reject_controls_and_line_breaks(self):
        for control in ("\n", "\r", "\t", "\x01", "\x7f", "\x85", "\u2028", "\u2029"):
            for entity, field in (("claims", "key"), ("claims", "subject"), ("claims", "predicate"),
                                  ("concepts", "key"), ("concepts", "title"), ("concepts", "description"),
                                  ("concepts", "resource"), ("concepts", "tags")):
                with self.subTest(control=repr(control), entity=entity, field=field):
                    state, _, payload = draft_state()
                    value = "safe" + control + "unexpected"
                    payload[entity][0][field] = [value] if field == "tags" else value
                    result = import_drafts(state, payload)
                    self.assertFalse(result["accepted"])
                    self.assertIn("single line", result["reason"])
                    self.assertFalse(state["claims"])

    def test_multiline_public_statement_is_retained_for_review(self):
        state, _, payload = draft_state()
        statement = "Training is offered on site.\n\n| Delivery | Audience |\n| --- | --- |\n| On site | Teams |"
        payload["claims"][0]["public_statement"] = statement
        self.assertTrue(import_drafts(state, payload)["accepted"])
        claim = next(iter(state["claims"].values()))
        self.assertEqual(claim["public_statement"], statement)
        self.assertEqual(claim["evidence_review"], "pending")

    def test_substantive_content_headers_and_footers_survive(self):
        for landmark in ("main", "section", "article"):
            with self.subTest(landmark=landmark):
                raw = f"<html><body><header>Site navigation noise</header><{landmark}><header><h1>Gutter clearance</h1><p>We clear gutters within 48 hours of booking.</p></header><p>Only buildings up to two storeys are eligible.</p><footer>Workshop content ends here.</footer></{landmark}><footer>Site legal footer noise</footer></body></html>"
                captured = extract_html(raw, "https://example.com/services/gutters/")
                self.assertIn("# Gutter clearance", captured["evidence"]["body"])
                self.assertIn("We clear gutters within 48 hours of booking.", captured["evidence"]["body"])
                self.assertIn("Only buildings up to two storeys", captured["evidence"]["body"])
                self.assertIn("Workshop content ends here.", captured["evidence"]["body"])
                self.assertNotIn("Site navigation noise", captured["evidence"]["body"])
                self.assertNotIn("Site legal footer noise", captured["evidence"]["body"])

    def test_draft_never_creates_review(self):
        state, _, payload = draft_state()
        self.assertTrue(import_drafts(state, payload)["accepted"])
        claim = next(iter(state["claims"].values()))
        self.assertEqual(claim["evidence_review"], "pending")
        self.assertEqual(claim["publication"], "pending")
        self.assertEqual(claim["reviews"], [])

    def test_draft_cannot_inject_approval(self):
        state, _, payload = draft_state()
        payload["claims"][0]["evidence_review"] = "accepted"
        self.assertFalse(import_drafts(state, payload)["accepted"])
        self.assertFalse(state["claims"])

    def test_quote_misattribution_and_fabrication(self):
        state, _, payload = draft_state()
        payload["claims"][0]["evidence"][0]["quote"] = "Training is free."
        self.assertFalse(import_drafts(state, payload)["accepted"])
        self.assertFalse(state["claims"])

    def test_quote_existence_is_not_semantic_approval(self):
        state, _, payload = draft_state()
        payload["claims"][0]["public_statement"] = "Training is not offered on site."
        self.assertTrue(import_drafts(state, payload)["accepted"])
        self.assertEqual(next(iter(state["claims"].values()))["evidence_review"], "pending")

    def test_changed_wording_invalidates_reviews(self):
        state, _, payload = draft_state()
        import_drafts(state, payload)
        old = next(iter(state["claims"].values()))
        old_revision = old["revision"]
        apply_review(state, "claim", old["id"], old_revision, {"evidence_review": "accepted"}, "human:agency-reviewer", utc_now(), "chat", "review-1")
        payload["claims"][0]["public_statement"] = "On-site training is offered."
        import_drafts(state, payload)
        claim = state["claims"][old["id"]]
        self.assertNotEqual(old_revision, claim["revision"])
        self.assertEqual(claim["evidence_review"], "pending")
        self.assertEqual(next(iter(state["concepts"].values()))["review"]["decision"], "pending")
        with self.assertRaises(ValueError):
            apply_review(state, "claim", old["id"], old_revision, {"publication": "approved"}, "human:agency-reviewer", utc_now(), "chat", "old-revision")

    def test_unchanged_import_preserves_reviews(self):
        state, _, payload = draft_state()
        import_drafts(state, payload)
        claim = next(iter(state["claims"].values()))
        apply_review(state, "claim", claim["id"], claim["revision"], {"evidence_review": "accepted"}, "human:agency-reviewer", utc_now(), "chat", "review-1")
        import_drafts(state, payload)
        self.assertEqual(next(iter(state["claims"].values()))["evidence_review"], "accepted")

    def test_risk_policy_and_confirmation_distinct(self):
        state, source, payload = draft_state()
        payload["claims"][0].update(public_statement="The fee is $100 per team.", predicate="fee", value=100)
        payload["claims"][0]["evidence"][0]["quote"] = "The fee is $100 per team."
        import_drafts(state, payload)
        claim = next(iter(state["claims"].values()))
        self.assertEqual(claim["business_confirmation"], "required")
        apply_review(state, "claim", claim["id"], claim["revision"], {"evidence_review": "accepted"}, "human:agency-reviewer", utc_now(), "chat", "risk-review")
        self.assertEqual(claim["business_confirmation"], "required")
        self.assertEqual(claim["publication"], "pending")

    def test_stock_photography_does_not_imply_live_inventory(self):
        state, source, payload = draft_state()
        raw = "<html><body><main><p>We provide stock photography for business publications.</p></main></body></html>"
        add_version(source, raw, extract_html(raw), "sf-original", state["config"])
        claim = payload["claims"][0]
        claim.update(subject="Stock photography", value="business publications", public_statement="We provide stock photography for business publications.")
        claim["evidence"][0].update(version_id=source["active_version"], quote=claim["public_statement"])
        self.assertTrue(import_drafts(state, payload)["accepted"])
        self.assertNotIn("volatile", next(iter(state["claims"].values()))["risk_categories"])

    def test_repair_bound_stops_invalid_batches(self):
        state, _, payload = draft_state()
        payload["claims"][0]["evidence"][0]["quote"] = "Missing quote"
        for _ in range(3):
            self.assertFalse(import_drafts(state, payload)["accepted"])
        with self.assertRaises(ValueError):
            import_drafts(state, payload)

    def test_anonymous_import_is_not_human_review(self):
        state, _, payload = draft_state()
        import_drafts(state, payload)
        claim = next(iter(state["claims"].values()))
        with self.assertRaises(ValueError):
            apply_review(state, "claim", claim["id"], claim["revision"], {"evidence_review": "accepted"}, "model:reviewer", utc_now(), "", "")

    def test_known_failed_source_cannot_be_waived(self):
        state, source, _ = draft_state()
        source["availability"] = "failed"
        with self.assertRaises(ValueError):
            accept_availability(state, [source["id"]], "human:client-owner", utc_now(), "offline acceptance", "chat", "decision")

    def test_wrong_owner_cannot_accept_availability(self):
        state, source, _ = draft_state()
        with self.assertRaises(ValueError):
            accept_availability(state, [source["id"]], "human:agency-reviewer", utc_now(), "offline acceptance", "chat", "decision")

    def test_workbook_roundtrip_reordered_rows_safe_text(self):
        state, _, payload = draft_state()
        payload["concepts"][0]["title"] = '=HYPERLINK("https://evil.invalid","click")'
        import_drafts(state, payload)
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "review.xlsx"
            export_review(state, path)
            book = load_workbook(path)
            self.assertEqual(book["Concepts"]["D2"].data_type, "s")
            book["Claims"]["G2"] = "accepted"
            book.save(path)
            import_review(state, path, "human:agency-reviewer", utc_now(), "returned-workbook")
            self.assertEqual(next(iter(state["claims"].values()))["evidence_review"], "accepted")
            self.assertEqual(next(iter(state["claims"].values()))["publication"], "pending")

    def test_workbook_material_edit_is_rejected_atomically(self):
        state, _, payload = draft_state()
        import_drafts(state, payload)
        before = copy.deepcopy(state)
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "review.xlsx"
            export_review(state, path)
            book = load_workbook(path)
            book["Claims"]["C2"] = "Training is free."
            book["Claims"]["G2"] = "accepted"
            book.save(path)
            with self.assertRaises(ValueError):
                import_review(state, path, "human:agency-reviewer", utc_now(), "returned-workbook")
        self.assertEqual(state, before)

    def test_formula_import_cannot_execute(self):
        state, _, payload = draft_state()
        import_drafts(state, payload)
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "review.xlsx"
            export_review(state, path)
            book = load_workbook(path)
            book["Claims"]["G2"] = '=WEBSERVICE("https://evil.invalid")'
            book.save(path)
            with self.assertRaises(ValueError):
                import_review(state, path, "human:agency-reviewer", utc_now(), "returned-workbook")

    def test_classification_updates_coverage_and_rejects_stale_override(self):
        state, source, payload = draft_state()
        import_drafts(state, payload)
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "review.xlsx"
            export_review(state, path)
            book = load_workbook(path)
            book["Classification"]["H2"] = "excluded"
            book["Classification"]["I2"] = "Outside selected scope"
            book.save(path)
            import_review(state, path, "human:agency-reviewer", utc_now(), "returned-workbook")
            self.assertEqual(state["inventory"][0]["state"], "excluded")
            book["Classification"]["H2"] = "included"
            book.save(path)
            with self.assertRaises(ValueError):
                import_review(state, path, "human:agency-reviewer", utc_now(), "stale-workbook")

    def test_agency_review_cannot_impersonate_client_confirmation(self):
        state, _, payload = draft_state()
        payload["claims"][0]["risk_categories"] = ["credential"]
        import_drafts(state, payload)
        claim = next(iter(state["claims"].values()))
        with self.assertRaises(ValueError):
            apply_review(state, "claim", claim["id"], claim["revision"], {"business_confirmation": "confirmed"}, "human:agency-reviewer", utc_now(), "chat", "not-client-confirmation")

    def test_client_source_column_is_bound_to_current_citation(self):
        state, source, payload = draft_state()
        payload["claims"][0]["risk_categories"] = ["credential"]
        import_drafts(state, payload)
        source = state["sources"][source["id"]]
        with tempfile.TemporaryDirectory() as temp:
            path, client = Path(temp) / "review.xlsx", Path(temp) / "client.xlsx"
            export_review(state, path, client)
            book = load_workbook(client)
            book.active["E2"] = "yes"
            book.active["G2"] = "Actual applicability scope"
            book.save(client)
            source["citation_url"] = "https://example.com/other/"
            with self.assertRaises(ValueError):
                import_review(state, client, "human:client-owner", utc_now(), "returned-workbook", client=True)

    def test_private_correction_does_not_replace_public_claim(self):
        state, _, payload = draft_state()
        payload["claims"][0]["risk_categories"] = ["disputed"]
        import_drafts(state, payload)
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "review.xlsx"
            client = Path(temp) / "client.xlsx"
            export_review(state, path, client)
            book = load_workbook(client)
            book.active["E2"] = "no"
            book.active["F2"] = "Training is available online only."
            book.active["G2"] = "Training offerings"
            book.save(client)
            import_review(state, client, "human:client-owner", utc_now(), "returned-workbook", client=True)
        claim = next(iter(state["claims"].values()))
        self.assertEqual(claim["public_statement"], "Training is offered on site.")
        self.assertEqual(claim["conflict"], "open")
        self.assertEqual(claim["publication"], "withheld")

    def test_refresh_invalidates_affected_evidence(self):
        state, source, payload = draft_state()
        import_drafts(state, payload)
        claim = next(iter(state["claims"].values()))
        apply_review(state, "claim", claim["id"], claim["revision"], {"evidence_review": "accepted"}, "human:agency-reviewer", utc_now(), "chat", "review")
        changed = copy.deepcopy(source)
        raw = "<html><body><main><p>Training moved online.</p></main></body></html>"
        add_version(changed, raw, extract_html(raw), "sf-original", state["config"])
        merge_import(state, {"sources": {source["id"]: changed}, "inventory": [], "warnings": []})
        self.assertEqual(claim["evidence_review"], "pending")
        self.assertIn(source["active_version"], state["sources"][source["id"]]["versions"])

    def test_unchanged_reimport_preserves_analyst_history_and_scope(self):
        state, source, _ = draft_state()
        source.update(scope="excluded", scope_reason="Analyst selected scope", alias_of="retained-alias-target", original_source_key=source["id"])
        source["review_history"] = [{"by": "human:agency-reviewer", "reference": "original-review"}]
        imported = copy.deepcopy(source)
        for key in ("alias_of", "original_source_key", "review_history"):
            imported.pop(key)
        imported.update(scope="included", scope_reason="Default suggestion")
        merge_import(state, {"sources": {source["id"]: imported}, "inventory": [], "warnings": []})
        current = state["sources"][source["id"]]
        self.assertEqual(current["alias_of"], "retained-alias-target")
        self.assertEqual(current["review_history"][0]["reference"], "original-review")
        self.assertEqual(current["scope"], "excluded")
        self.assertEqual(state["inventory"][0]["state"], "excluded")

    def test_reimport_cannot_erase_known_http_failure(self):
        for status in (401, 403, 404, 410, 429, 500, 502):
            state, source, _ = draft_state()
            source["captured_status"] = status
            incoming = copy.deepcopy(source)
            incoming["captured_status"] = None
            merge_import(state, {"sources": {source["id"]: incoming}, "inventory": [], "warnings": []})
            current = state["sources"][source["id"]]
            self.assertEqual(current["captured_status"], status)
            self.assertIn(current["availability"], ("failed", "nonpublic"))
            with self.assertRaises(ValueError):
                accept_availability(state, [source["id"]], "human:client-owner", utc_now(), "Cannot waive known failure", "chat", "fixture")

    def test_observed_availability_recovery_is_distinct_from_capture(self):
        state, source, payload = draft_state()
        import_drafts(state, payload)
        source = state["sources"][source["id"]]
        source.update(captured_status=500, availability="failed")
        raw = source["versions"][source["active_version"]]["raw_content"].encode()
        check_availability(state, [source["id"]], fetch=lambda url: {"body": raw, "status": 200, "url": url})
        self.assertEqual(source["captured_status"], 500)
        self.assertEqual(source["availability"], "live")
        self.assertTrue(source["availability_check"]["supersedes_captured_failure"])
        self.assertFalse(source["availability_check"]["requires_source_refresh"])

    def test_changed_live_body_cannot_reuse_old_evidence_reviews(self):
        state, source, payload = draft_state()
        import_drafts(state, payload)
        source = state["sources"][source["id"]]
        claim = next(iter(state["claims"].values()))
        apply_review(state, "claim", claim["id"], claim["revision"], {"evidence_review": "accepted"}, "human:agency-reviewer", utc_now(), "chat", "old-source")
        raw = b"<html><body><main>Training is offered on site. A new exception applies to consumer clients. The fee is $100 per team.</main></body></html>"
        check_availability(state, [source["id"]], fetch=lambda url: {"body": raw, "status": 200, "url": url})
        self.assertTrue(source["availability_check"]["requires_source_refresh"])
        self.assertEqual(claim["evidence_review"], "pending")
        self.assertEqual(claim["reviews"], [])

    def test_packets_cover_every_selected_page_without_model_calls(self):
        state, _, _ = draft_state()
        with tempfile.TemporaryDirectory() as temp:
            report = packets(state, temp, max_chars=1000)
            self.assertGreater(report["selected_evidence_characters"], 0)
            self.assertEqual(report["actual_chat_usage"], "unavailable")
            self.assertTrue(report["packet_files"])

    def test_url_variants_fragments_preserved(self):
        self.assertEqual(normalize_url("HTTPS://EXAMPLE.COM:443/%7eUser/?x=1&x=2#/a"), "https://example.com/~User/?x=1&x=2#/a")
        self.assertNotEqual(normalize_url("https://example.com/A"), normalize_url("https://example.com/a"))
        self.assertNotEqual(normalize_url("https://example.com/#/a"), normalize_url("https://example.com/#/b"))

    def test_alias_needs_explicit_review_and_retains_capture_history(self):
        state, source, _ = draft_state()
        target = new_source("https://example.com/training", state["config"])
        active = source["versions"][source["active_version"]]
        target["versions"] = copy.deepcopy(source["versions"])
        target["active_version"] = active["id"]
        state["sources"][target["id"]] = target
        self.assertNotEqual(source["source_key"], target["source_key"])
        approve_alias(state, source["id"], target["id"], "human:agency-reviewer", utc_now(), "chat", "alias-evidence-review", "Confirmed identical capture aliases")
        self.assertEqual(source["source_key"], target["source_key"])
        self.assertEqual(source["original_source_key"], source["id"])
        self.assertIn(active["id"], source["versions"])
        self.assertNotEqual(source["citation_url"], target["citation_url"])
        with self.assertRaises(ValueError):
            approve_alias(state, target["id"], source["id"], "human:agency-reviewer", utc_now(), "chat", "cycle", "No silent alias cycle")

    def test_divergent_hash_views_are_not_alias_merged(self):
        state, source, _ = draft_state()
        target = new_source("https://example.com/#/different-view", state["config"])
        raw = "<html><body><main>Different view content and unique service qualifications.</main></body></html>"
        add_version(target, raw, extract_html(raw), "sf-original", state["config"])
        state["sources"][target["id"]] = target
        with self.assertRaises(ValueError):
            approve_alias(state, source["id"], target["id"], "human:agency-reviewer", utc_now(), "chat", "not-identical", "Views differ")

    def test_ambiguous_filename_candidates_do_not_share_source_keys(self):
        from okf_bundle.ingest import ingest_sf
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "ambiguous.zip"
            with zipfile.ZipFile(path, "w") as archive:
                for prefix in ("a", "b"):
                    archive.writestr(prefix + "/original_https_example.com_uncertain_.html", "<html><body><main>Distinct unknown capture identity needs human attribution before publication.</main></body></html>")
            result = ingest_sf(path, {"domains": ["example.com"], "locale": "en"})
            self.assertEqual(len(result["sources"]), 2)
            self.assertEqual(len({s["source_key"] for s in result["sources"].values()}), 2)
            self.assertTrue(all(not s["identity_confirmed"] for s in result["sources"].values()))

    def test_state_rejects_concurrent_writer_and_preserves_history(self):
        state, _, _ = draft_state()
        with tempfile.TemporaryDirectory() as temp:
            store = StateStore(temp)
            with store.transaction():
                first = store.save(state)
                with self.assertRaises(RuntimeError):
                    with StateStore(temp).transaction():
                        pass
                state["config"]["model"] = "actual-model"
                second = store.save(state)
            self.assertNotEqual(first, second)
            self.assertEqual(len(list((store.private / "snapshots").glob("*.json"))), 2)
            self.assertEqual(store.load()["config"]["model"], "actual-model")

    def test_interrupted_pointer_commit_keeps_previous_snapshot(self):
        state, _, _ = draft_state()
        with tempfile.TemporaryDirectory() as temp:
            store = StateStore(temp)
            with store.transaction():
                store.save(state)
                state["config"]["model"] = "new-model"
                import okf_bundle.state as module
                original = module.atomic_json
                def interrupt(path, value):
                    if Path(path).name == "current.json":
                        raise OSError("fixture interrupted before pointer swap")
                    original(path, value)
                with patch.object(module, "atomic_json", interrupt):
                    with self.assertRaises(OSError):
                        store.save(state)
            self.assertNotEqual(store.load()["config"]["model"], "new-model")

    def test_live_lock_cannot_be_recovered(self):
        state, _, _ = draft_state()
        with tempfile.TemporaryDirectory() as temp:
            store = StateStore(temp)
            with store.transaction():
                store.save(state)
                with self.assertRaises(RuntimeError):
                    store.recover_lock("test live owner")

    def test_snapshot_tampering_fails_integrity_check(self):
        state, _, _ = draft_state()
        with tempfile.TemporaryDirectory() as temp:
            store = StateStore(temp)
            with store.transaction():
                key = store.save(state)
            path = store.private / "snapshots" / (key + ".json")
            path.write_text("{}", encoding="utf-8")
            with self.assertRaises(ValueError):
                store.load()


if __name__ == "__main__":
    unittest.main()
