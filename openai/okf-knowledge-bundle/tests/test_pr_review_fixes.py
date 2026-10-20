import copy
import json
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from okf_bundle.common import digest
from okf_bundle.ingest import discover_sitemap
from okf_bundle.workbook import export_review, import_review
from okf_bundle.workflow import packets
from test_review_state import draft_state


class PullRequestFixTests(unittest.TestCase):
    def test_stale_classification_cannot_approve_new_restrictions(self):
        changes = {
            'http_robots': 'noindex', 'x_robots_tag': 'noindex', 'robots': 'noindex',
            'csv_indexability': 'Non-Indexable', 'reader_warning_review_required': True,
            'managed_signal_review_required': True,
            'managed_capture_signals': [{'robots': 'noindex', 'review_required': True}],
            'warnings': ['New provider warning requires review'],
        }
        for field, value in changes.items():
            with self.subTest(field=field), tempfile.TemporaryDirectory() as temp:
                state, source, _ = draft_state()
                path = Path(temp) / 'review.xlsx'
                export_review(state, path)
                book = load_workbook(path)
                book['Classification']['H2'] = 'included'
                book['Classification']['I2'] = 'Synthetic decision before restriction'
                book.save(path); book.close()
                active = source['active_version']
                source[field] = value; source['inclusion_decision'] = None
                before = digest(state)
                with self.assertRaisesRegex(ValueError, 'Stale/modified Classification'):
                    import_review(state, path, 'human:fixture-reviewer', datetime.now(timezone.utc).isoformat(), 'synthetic-test', 'old-workbook')
                self.assertEqual(before, digest(state))
                self.assertEqual(active, source['active_version'])
                export_review(state, Path(temp) / 'fresh.xlsx')
                fresh = load_workbook(Path(temp) / 'fresh.xlsx', read_only=True)
                self.assertIn(field, fresh['Sources Coverage']['I2'].value)
                fresh.close()

    def test_regenerated_packets_remove_old_versions_chunks_and_exclusions(self):
        state, source, _ = draft_state()
        version = source['versions'][source['active_version']]
        version['evidence']['body'] = 'Training is offered on site. ' * 120
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            old = packets(state, folder, max_chars=1000)
            self.assertGreater(len(old['packet_files']), 1)
            unrelated = folder / 'operator-note.json'; unrelated.write_text('{}')
            source['active_version'] = 'ver-' + 'a' * 16
            source['versions'][source['active_version']] = copy.deepcopy(version)
            source['versions'][source['active_version']]['evidence']['body'] = 'Training is offered on site.'
            current = packets(state, folder, max_chars=1000)
            self.assertFalse(set(old['packet_files']) & set(current['packet_files']))
            self.assertEqual(set(current['packet_files']) | {'packet-manifest.json', unrelated.name}, {p.name for p in folder.iterdir()})
            source['scope'] = 'excluded'
            empty = packets(state, folder, max_chars=1000)
            self.assertEqual([], empty['packet_files'])
            self.assertEqual({'packet-manifest.json', unrelated.name}, {p.name for p in folder.iterdir()})
            self.assertEqual([], json.loads((folder / 'packet-manifest.json').read_text())['packet_files'])

    def test_sitemap_limit_counts_duplicate_and_missing_url_occurrences(self):
        for repeated in ['<url><loc>https://example.com/a/</loc></url>', '<url/>']:
            with self.subTest(entry=repeated), tempfile.TemporaryDirectory() as temp:
                path = Path(temp) / 'site.xml'; path.write_text('<urlset>' + repeated * 5 + '</urlset>')
                result = discover_sitemap(path, {'domains': ['example.com'], 'limits': {'urls': 3}})
                self.assertEqual(3, len(result['inventory']))
                self.assertTrue(result['sitemap_summary']['partial'])


if __name__ == '__main__':
    unittest.main()
