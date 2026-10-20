"""Portable private analyst/client workbooks; untrusted cells are always text."""
import copy
import json
import zipfile
import math
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from .common import digest
from .review import apply_review, reviewer_event


def _source_restrictions(source):
    signals = {key: source.get(key) for key in (
        "robots", "http_robots", "x_robots_tag", "csv_indexability", "warnings",
        "reader_warning_review_required", "managed_signal_review_required",
        "managed_capture_signals", "robots_disallowed", "csv_metadata_review_required",
        "captured_status", "availability",
    )}
    version = source.get("versions", {}).get(source.get("active_version"), {})
    signals["active_version_warnings"] = version.get("warnings", [])
    signals["active_version_quality"] = version.get("quality")
    return signals


def source_revision(source):
    material = {k: source.get(k) for k in ("id", "source_key", "alias_of", "normalized_url", "citation_url", "active_version", "identity_basis", "canonical", "og_url", "scope", "scope_reason", "classification")}
    material["restrictions"] = _source_restrictions(source)
    return digest(material)


def _text(value):
    if value is None:
        return ""
    result = json.dumps(value, ensure_ascii=False, sort_keys=True) if isinstance(value, (dict, list)) else str(value)
    if len(result) > 32767:
        raise ValueError("Review cell exceeds Excel's text limit; split the claim/evidence instead of truncating")
    return result


def _sheet(book, name, headers, rows, editable=(), choices=None):
    sheet = book.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        sheet.append([_text(value) for value in row])
    for cells in sheet:
        for cell in cells:
            cell.data_type = "s"  # =,+,-,@ payloads are text, never formulas.
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.protection = Protection(locked=cell.column not in editable or cell.row == 1)
            if cell.row == 1:
                cell.font = Font(bold=True, color="FFFFFF", size=11)
                cell.fill = PatternFill("solid", fgColor="173D4A")
            elif cell.column in editable:
                cell.fill = PatternFill("solid", fgColor="E7F2FA")
                cell.font = Font(color="145AA3", size=11)
            elif cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F1F5F6")
    for index, header in enumerate(headers, 1):
        from openpyxl.utils import get_column_letter
        width = 24 if "revision" not in header.lower() and "id" != header.lower() else 20
        if any(word in header.lower() for word in ("statement", "description", "evidence", "reason", "quote", "title")):
            width = 58
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.zoomScale = 85
    sheet.row_dimensions[1].height = 32
    for row in range(2, sheet.max_row + 1):
        lines = max(sum(max(1, math.ceil(len(part) / max(10, sheet.column_dimensions[cell.column_letter].width - 2))) for part in str(cell.value or "").split("\n")) for cell in sheet[row])
        sheet.row_dimensions[row].height = min(409, max(45, lines * 15 + 8))
    sheet.protection.sheet = True  # Convenience only; importer enforces IDs/revisions.
    sheet.protection.autoFilter = False
    if choices:
        for column, values in choices.items():
            validation = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
            validation.errorTitle = "Choose a listed review decision"
            validation.error = "Use one of the choices in this cell."
            validation.showErrorMessage = True
            sheet.add_data_validation(validation)
            from openpyxl.utils import get_column_letter
            validation.add(f"{get_column_letter(column)}2:{get_column_letter(column)}{max(2, sheet.max_row)}")
    return sheet


def _evidence_text(claim, sources):
    return "\n\n".join((sources[e["source_id"]].get("citation_url") or "Source identity needs review") + "\n" + e["source_id"] + " / " + e["version_id"] + " / " + e["locator"] + "\n" + e["quote"] for e in claim["evidence"])


def export_review(state, path, client_path=None):
    book = Workbook()
    book.remove(book.active)
    _sheet(book, "Release Summary", ["Item", "Value"], [
        ["Run ID", state["run_id"]], ["Configuration fingerprint", state["config_hash"]],
        ["Stage", "Private draft/review — not a public release"],
        ["Sources", len(state["sources"])], ["Claims", len(state["claims"])], ["Concepts", len(state["concepts"])],
        ["Evidence review", "Accept only after checking the statement, qualifiers and exact source passage."],
        ["Publication", "Claim decisions do not replace approval of the frozen release manifest."],
        ["Corrections", "Send corrections through the assistant; new wording requires public supporting evidence and a new revision."],
        ["Long passages", "Rows expand up to Excel's limit. Select a long cell and expand the formula bar to inspect its complete retained text before deciding."],
    ])
    sources = state["sources"]
    _sheet(book, "Sources Coverage", ["ID", "Revision", "Captured identity URL", "Canonical", "og:url", "Identity basis", "Availability", "Captured HTTP status", "Warnings"], [
        [s["id"], source_revision(s), s.get("requested_url"), s.get("canonical"), s.get("og_url"), s.get("identity_basis"), s.get("availability", "unknown"), s.get("captured_status"), _source_restrictions(s)] for s in sources.values()
    ])
    _sheet(book, "Classification", ["ID", "Revision", "URL", "Suggested purpose", "Signals", "Scope", "Override purpose", "Scope decision", "Reason"], [
        [s["id"], source_revision(s), s.get("requested_url"), s.get("classification", {}).get("purpose"), s.get("classification", {}).get("signals", []), s.get("scope", "pending"), "", "", ""] for s in sources.values()
    ], editable=(7, 8, 9), choices={8: ["included", "excluded", "pending"]})
    claim_rows = []
    for claim in state["claims"].values():
        claim_rows.append([claim["id"], claim["revision"], claim["public_statement"], claim["qualifiers"], _evidence_text(claim, sources), claim["risk_categories"], claim["evidence_review"], claim["business_confirmation"], claim["publication"], claim["conflict"], ""])
    _sheet(book, "Claims", ["ID", "Revision", "Public statement", "Qualifiers", "Evidence passages", "Confirmation risks", "Evidence review", "Business confirmation", "Publication", "Conflict", "Conflict resolution reason"], claim_rows, editable=(7, 8, 9, 10, 11), choices={7: ["pending", "accepted", "rejected"], 8: ["not_required", "required", "confirmed", "disputed"], 9: ["pending", "approved", "withheld"], 10: ["none", "open", "resolved"]})
    _sheet(book, "Conflicts", ["Claim ID", "Revision", "Public statement", "Conflict", "Evidence"], [[c["id"], c["revision"], c["public_statement"], c["conflict"], c["evidence"]] for c in state["claims"].values() if c["conflict"] != "none"])
    _sheet(book, "Concepts", ["ID", "Revision", "Type", "Title", "Description", "Resource", "Claim IDs", "Decision"], [[c["id"], c["revision"], c["type"], c["title"], c["description"], c.get("resource"), c["claim_ids"], c.get("review", {}).get("decision", "pending")] for c in state["concepts"].values()], editable=(8,), choices={8: ["pending", "accepted", "rejected"]})
    _sheet(book, "Changes", ["Kind", "ID", "Reason"], [["Pending draft", r.get("batch_id"), r.get("reason")] for r in state.get("draft_failures", [])] + [["Input", row.get("input_ref"), row.get("state", "") + ": " + row.get("reason", "")] for row in state["inventory"]])
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    if client_path:
        client = Workbook()
        client.remove(client.active)
        _sheet(client, "Client confirmation", ["ID", "Revision", "Public statement", "Public source", "Correct?", "Correction", "Reviewer scope"], [[c["id"], c["revision"], c["public_statement"], "; ".join(sources[e["source_id"]].get("citation_url") or "Identity needs review" for e in c["evidence"]), "", "", ""] for c in state["claims"].values() if c["business_confirmation"] != "not_required"], editable=(5, 6, 7), choices={5: ["yes", "no", "unsure"]})
        Path(client_path).parent.mkdir(parents=True, exist_ok=True)
        client.save(client_path)
    return {"review": str(path), "client_confirmation": str(client_path) if client_path else None}


def _rows(book, sheet, headers):
    if sheet not in book.sheetnames:
        raise ValueError("Missing review sheet: " + sheet)
    ws = book[sheet]
    values = list(ws.iter_rows(values_only=True))
    if not values or list(values[0]) != headers:
        raise ValueError("Changed review columns: " + sheet)
    seen = set()
    for row in values[1:]:
        if all(v is None for v in row):
            continue
        if row[0] in seen:
            raise ValueError("Duplicate review ID: " + str(row[0]))
        seen.add(row[0])
        yield ["" if value is None else str(value) for value in row]


def import_review(state, path, by, at, channel, reference=None, client=False):
    path = Path(path)
    if path.stat().st_size > 25 * 1024 * 1024:
        raise ValueError("Review workbook exceeds 25 MB")
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 2000 or sum(e.file_size for e in entries) > 100 * 1024 * 1024 or any(e.file_size > 25 * 1024 * 1024 for e in entries):
            raise ValueError("Review workbook exceeds expanded XML/member bounds")
    # Do not resolve external links or macros; formula cells are never evaluated.
    book = load_workbook(path, data_only=False, keep_links=False, read_only=True)
    try:
        return _apply_workbook(state, path, book, by, at, channel, reference, client)
    finally:
        book.close()


def _apply_workbook(state, path, book, by, at, channel, reference, client):
    for ws in book:
        if ws.max_row > 50001 or ws.max_column > 30:
            book.close()
            raise ValueError("Review worksheet exceeds row/column bounds")
        for row in ws:
            for cell in row:
                if cell.data_type == "f":
                    raise ValueError(f"Formula cell is not accepted: {ws.title}!{cell.coordinate}")
    file_hash = digest(path.read_bytes())
    reference = (reference + "; " if reference else "") + "sha256:" + file_hash
    reviewer_event(by, at, channel, reference, "import")
    working = copy.deepcopy(state)
    count = 0
    if client:
        for row in _rows(book, "Client confirmation", ["ID", "Revision", "Public statement", "Public source", "Correct?", "Correction", "Reviewer scope"]):
            cid, revision, statement, source, correct, correction, scope = row
            claim = working["claims"].get(cid)
            if not claim or revision != claim["revision"] or statement != claim["public_statement"]:
                raise ValueError(f"Stale/modified Client confirmation row for {cid}")
            expected_source = "; ".join(working["sources"][e["source_id"]].get("citation_url") or "Identity needs review" for e in claim["evidence"])
            if source != expected_source:
                raise ValueError("Stale/modified public source in client confirmation row")
            if correct not in ("", "yes", "no", "unsure"):
                raise ValueError("Invalid client confirmation choice")
            if not correct and not correction:
                continue
            if not scope:
                raise ValueError("Client decision needs reviewer scope")
            decisions = {"business_confirmation": "confirmed" if correct == "yes" and not correction else "disputed"}
            if correction or correct == "no":
                decisions.update(conflict="open", publication="withheld")
                working.setdefault("private_corrections", []).append({"claim_id": cid, "revision": revision, "correction": correction, "by": by, "at": at, "scope": scope, "reference": reference, "status": "needs-public-support"})
            apply_review(working, "claim", cid, revision, decisions, by, at, channel, reference + "; scope:" + scope)
            count += 1
    else:
        for row in _rows(book, "Claims", ["ID", "Revision", "Public statement", "Qualifiers", "Evidence passages", "Confirmation risks", "Evidence review", "Business confirmation", "Publication", "Conflict", "Conflict resolution reason"]):
            cid, revision, statement, qualifiers, evidence, risks, *tail = row
            values, resolution_reason = tail[:4], tail[4]
            claim = working["claims"].get(cid)
            if not claim or revision != claim["revision"] or [statement, qualifiers, evidence, risks] != [_text(claim["public_statement"]), _text(claim["qualifiers"]), _evidence_text(claim, working["sources"]), _text(claim["risk_categories"])]:
                raise ValueError(f"Stale/modified Claims row for {cid}; resubmit changed wording as a new draft")
            names = ("evidence_review", "business_confirmation", "publication", "conflict")
            changes = {key: val for key, val in zip(names, values) if val != claim[key]}
            if changes:
                apply_review(working, "claim", cid, revision, changes, by, at, channel, reference, reason=resolution_reason or None)
                count += 1
        for row in _rows(book, "Concepts", ["ID", "Revision", "Type", "Title", "Description", "Resource", "Claim IDs", "Decision"]):
            cid, revision, *values, decision = row
            concept = working["concepts"].get(cid)
            if not concept or revision != concept["revision"] or values != [_text(concept.get(k)) for k in ("type", "title", "description", "resource", "claim_ids")]:
                raise ValueError(f"Stale/modified Concepts row for {cid}")
            if decision != concept.get("review", {}).get("decision", "pending"):
                apply_review(working, "concept", cid, revision, {"decision": decision}, by, at, channel, reference)
                count += 1
        for row in _rows(book, "Classification", ["ID", "Revision", "URL", "Suggested purpose", "Signals", "Scope", "Override purpose", "Scope decision", "Reason"]):
            sid, revision, url, purpose, signals, original_scope, override, scope, reason = row
            source = working["sources"].get(sid)
            if not source or revision != source_revision(source) or url != _text(source.get("requested_url")):
                raise ValueError(f"Stale/modified Classification row for {sid}")
            if not override and not scope:
                continue
            if not reason or scope not in ("", "included", "excluded", "pending"):
                raise ValueError("Classification override needs a reason and valid scope")
            if override:
                source.setdefault("classification", {})["analyst_override"] = override
            if scope:
                source["scope"] = scope
                source["scope_reason"] = reason
                source["inclusion_decision"] = {"by": by, "at": at, "channel": channel, "reference": reference, "version_id": source.get("active_version"), "decision": scope, "reason": reason}
                for item in working["inventory"]:
                    if item.get("source_id") == sid and item.get("state") != "duplicate":
                        item.update(state=scope, reason=reason)
            source.setdefault("review_history", []).append({"by": by, "at": at, "reference": reference, "channel": channel, "reason": reason, "revision": revision})
            working["release_approvals"] = {}
            count += 1
    state.update(working)
    return {"changed_rows": count, "import_sha256": file_hash}
