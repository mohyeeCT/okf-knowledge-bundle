"""Regression checks for the independently reviewed analyst workflow defects."""
import copy
from datetime import timedelta
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from openpyxl import load_workbook
from okf_bundle.cli import execute, parser
from okf_bundle.common import parse_time, utc_now
from okf_bundle.ingest import ingest_sf, discover_csv
from okf_bundle.review import import_drafts
from okf_bundle.state import StateStore, new_state
from okf_bundle.workbook import export_review, import_review
from okf_bundle.workflow import merge_import
from test_cli import HTML
from test_export import fixture_state
from test_review_state import draft_state


class ReviewWorkflowFixTests(unittest.TestCase):
    def test_reimport_updates_inventory_without_losing_distinct_rows(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            archive = root / 'site.zip'
            with zipfile.ZipFile(archive, 'w') as supplied:
                supplied.writestr('original_https_example.com_training_.html', HTML)
                supplied.writestr('original_https_example.com_other_.html', HTML.replace('/training/', '/other/'))
            state = new_state({'domains': ['example.com']})
            merge_import(state, ingest_sf(archive, state['config']))
            before = copy.deepcopy(state['inventory'])
            merge_import(state, ingest_sf(archive, state['config']))
            self.assertEqual(state['inventory'], before)
            urls = root / 'urls.csv'
            urls.write_text('URL\nhttps://example.com/a/\nhttps://example.com/a/\n')
            another = new_state(state['config'])
            result = discover_csv(urls, another['config'])
            merge_import(another, result)
            merge_import(another, result)
            self.assertEqual(len(another['inventory']), 2)
            self.assertEqual([row['state'] for row in another['inventory']], ['pending', 'duplicate'])

    def test_reimport_cannot_erase_redirect_or_unresolved_csv_observation(self):
        state, source, _ = draft_state()
        source['redirect_review_required'] = {'to_url': 'https://example.com/new/'}
        source['csv_metadata_review_required'] = True
        source['input_metadata'] = [{'values': {'Status Code': '404'}}]
        incoming = copy.deepcopy(source)
        for key in ('redirect_review_required', 'csv_metadata_review_required', 'input_metadata'):
            incoming.pop(key)
        merge_import(state, {'sources': {source['id']: incoming}, 'inventory': [], 'warnings': []})
        current = state['sources'][source['id']]
        self.assertTrue(current['csv_metadata_review_required'])
        self.assertEqual(current['input_metadata'][0]['values']['Status Code'], '404')
        self.assertIn('redirect_review_required', current)

    def test_collect_reports_actual_stop_summary(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            run = root / 'run'
            config = root / 'config.json'
            config.write_text(json.dumps({'domains': ['example.com'], 'collection_authorized': True}))
            execute(parser().parse_args(['init', '--run', str(run), '--config', str(config)]))
            urls = root / 'urls.csv'
            urls.write_text('URL\nhttps://example.com/training/\nhttps://example.com/other/\n')
            execute(parser().parse_args(['discover-csv', '--run', str(run), '--csv', str(urls)]))
            def reader(url, settings):
                return {'provider_status': 200, 'data': {'html': HTML.replace('https://example.com/training/', url), 'url': url, 'usage': {'tokens': 10}}}
            with patch.dict(os.environ, {'JINA_API_KEY': 'fixture-no-network'}), patch('okf_bundle.network.jina_request', reader):
                result = execute(parser().parse_args(['collect', '--run', str(run), '--max-pages', '1']))
            self.assertEqual(result['collection']['attempts'], 1)
            self.assertEqual(result['collection']['pending'], 1)
            self.assertIn('ceiling', result['collection']['stopped_reason'])

    def test_release_approval_rejects_time_before_candidate_without_saving(self):
        with tempfile.TemporaryDirectory() as folder:
            store = StateStore(Path(folder) / 'run')
            with store.transaction():
                store.save(fixture_state())
            args = ['--run', str(store.root)]
            candidate = execute(parser().parse_args(['candidate', *args, '--release-id', 'time-check']))
            before = StateStore(store.root).load()
            earlier = (parse_time(utc_now()) - timedelta(days=1)).isoformat()
            with self.assertRaisesRegex(ValueError, 'cannot predate'):
                execute(parser().parse_args(['release-approve', *args, '--candidate', candidate['candidate'],
                    '--manifest-hash', candidate['manifest_hash'], '--by', 'human:client-owner', '--at', earlier,
                    '--channel', 'synthetic-fixture', '--reference', 'not-real-approval']))
            self.assertEqual(before, StateStore(store.root).load())

    def test_workbook_failure_names_claim_and_preserves_all_decisions(self):
        state, _, payload = draft_state()
        import_drafts(state, payload)
        cid = next(iter(state['claims']))
        before = copy.deepcopy(state)
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'review.xlsx'
            export_review(state, path)
            book = load_workbook(path)
            book['Claims']['C2'] = 'Changed factual statement'
            book['Claims']['G2'] = 'accepted'
            book.save(path)
            book.close()
            with self.assertRaisesRegex(ValueError, 'Claims row for ' + cid):
                import_review(state, path, 'human:analyst', utc_now(), 'returned-fixture')
        self.assertEqual(state, before)

    def test_cli_followup_candidate_keeps_verified_log_history(self):
        with tempfile.TemporaryDirectory() as folder:
            store = StateStore(Path(folder) / 'run')
            with store.transaction():
                store.save(fixture_state())
            run = ['--run', str(store.root)]
            first = execute(parser().parse_args(['candidate', *run, '--release-id', 'first']))
            second = execute(parser().parse_args(['candidate', *run, '--release-id', 'second',
                '--previous-candidate', first['candidate']]))
            log = (Path(second['candidate']) / 'okf/log.md').read_text(encoding='utf-8')
            self.assertIn('first', log)
            self.assertIn('second', log)

    def test_cli_comparison_acceptance_binds_current_response_and_capture(self):
        with tempfile.TemporaryDirectory() as folder:
            store = StateStore(Path(folder) / 'run')
            state = fixture_state()
            source = state['sources']['source-1']
            source['availability_check'] = {'at': utc_now(), 'version_id': source['active_version'],
                'response_hash': 'fixture-response-a', 'comparison_inconclusive': True,
                'requires_source_refresh': False, 'redirected': False, 'quotes_supported': True}
            with store.transaction():
                store.save(state)
            command = ['source-review', '--run', str(store.root), '--id', source['id'],
                '--url', source['citation_url'], '--scope', 'included', '--accept-comparison',
                '--reason', 'Synthetic human compared the displayed rendered and raw evidence.',
                '--by', 'human:analyst', '--at', utc_now(), '--channel', 'synthetic-fixture',
                '--reference', 'not-real-approval']
            execute(parser().parse_args(command))
            decision = StateStore(store.root).load()['sources'][source['id']]['availability_comparison_review']
            self.assertEqual(decision['version_id'], source['active_version'])
            self.assertEqual(decision['response_hash'], 'fixture-response-a')
            self.assertEqual(decision['decision'], 'accepted')
            with store.transaction():
                current = store.load()
                current['sources'][source['id']]['availability_check']['redirected'] = True
                store.save(current)
            with self.assertRaisesRegex(ValueError, 'No current eligible'):
                execute(parser().parse_args(command))

    def test_workbook_resolution_requires_explicit_human_reason(self):
        state, _, payload = draft_state()
        payload['claims'][0]['risk_categories'] = ['disputed']
        import_drafts(state, payload)
        cid = next(iter(state['claims']))
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'review.xlsx'
            export_review(state, path)
            book = load_workbook(path)
            book['Claims']['J2'] = 'resolved'
            book.save(path)
            with self.assertRaisesRegex(ValueError, 'reason'):
                import_review(state, path, 'human:analyst', utc_now(), 'returned-fixture')
            self.assertEqual(state['claims'][cid]['conflict'], 'open')
            book['Claims']['K2'] = 'Fictional reviewer checked both passages and their effective periods.'
            book.save(path)
            book.close()
            import_review(state, path, 'human:analyst', utc_now(), 'returned-fixture')
            self.assertEqual(state['claims'][cid]['conflict'], 'resolved')
            self.assertIn('both passages', state['claims'][cid]['reviews'][-1]['reason'])


if __name__ == '__main__':
    unittest.main()
